#!/usr/bin/env python3
"""
Markdown to DOCX Converter for Dissertation Report
====================================================
Converts FINAL_COMPILED_MIDSEM_REPORT.md to a professional DOCX document
with proper formatting, headings, tables, lists, and styling.
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_background(cell, fill):
    """Set cell background color."""
    shading_elm = OxmlElement("w:shd")
    shading_elm.set(qn("w:fill"), fill)
    cell._element.get_or_add_tcPr().append(shading_elm)


def parse_markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    """Convert Markdown to DOCX with formatting."""
    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    with md_path.open("r", encoding="utf-8") as fh:
        lines = fh.readlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # Heading 1
        if line.startswith("# "):
            text = line[2:].strip()
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.size = Pt(16)
            run.font.bold = True
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1

        # Heading 2
        elif line.startswith("## "):
            text = line[3:].strip()
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0, 51, 102)
            i += 1

        # Heading 3
        elif line.startswith("### "):
            text = line[4:].strip()
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.font.size = Pt(12)
            run.font.bold = True
            i += 1

        # Horizontal rule
        elif line.strip() == "---":
            doc.add_paragraph()
            i += 1

        # Table (markdown format)
        elif "|" in line and i + 1 < len(lines) and "|" in lines[i + 1]:
            table_lines = [line]
            i += 1
            # Skip separator line
            if "|" in lines[i] and "-" in lines[i]:
                i += 1
            # Collect table rows
            while i < len(lines) and "|" in lines[i]:
                table_lines.append(lines[i])
                i += 1

            # Parse table
            headers = [cell.strip() for cell in table_lines[0].split("|")[1:-1]]
            rows = []
            for table_line in table_lines[1:]:
                cells = [cell.strip() for cell in table_line.split("|")[1:-1]]
                if cells and len(cells) == len(headers):
                    rows.append(cells)

            # Create table
            if rows:
                table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
                table.style = "Light Grid Accent 1"

                # Header row
                header_cells = table.rows[0].cells
                for j, header in enumerate(headers):
                    header_cells[j].text = header
                    for paragraph in header_cells[j].paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                    set_cell_background(header_cells[j], "D3D3D3")

                # Data rows
                for row_idx, row_data in enumerate(rows, 1):
                    cells = table.rows[row_idx].cells
                    for col_idx, cell_data in enumerate(row_data):
                        cells[col_idx].text = cell_data

        # Unordered list
        elif line.strip().startswith("- "):
            text = line.strip()[2:].strip()
            text = format_inline_text(text)
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(text)
            i += 1

        # Numbered list
        elif line.strip() and re.match(r"^\d+\.", line.strip()):
            match = re.match(r"^(\d+)\.\s+(.*)", line.strip())
            if match:
                text = match.group(2).strip()
                text = format_inline_text(text)
                p = doc.add_paragraph(style="List Number")
                run = p.add_run(text)
            i += 1

        # Code block
        elif line.strip().startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip())
                i += 1
            if i < len(lines):
                i += 1  # Skip closing ```
            code_text = "\n".join(code_lines)
            p = doc.add_paragraph(code_text, style="No Spacing")
            for run in p.runs:
                run.font.name = "Courier New"
                run.font.size = Pt(10)
            i += 1

        # Regular paragraph
        else:
            text = line.strip()
            if text:
                text = format_inline_text(text)
                p = doc.add_paragraph(text)
            i += 1

    # Set margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    doc.save(docx_path)
    print(f"✓ Converted {md_path.name} → {docx_path.name}")


def format_inline_text(text: str) -> str:
    """Placeholder for inline formatting (bold, italic).
    Note: python-docx requires manual run creation for rich inline formatting."""
    return text


def main():
    base_path = Path(__file__).parent.parent
    md_path = base_path / "FINAL_COMPILED_MIDSEM_REPORT.md"
    docx_path = base_path / "FINAL_COMPILED_MIDSEM_REPORT.docx"

    if not md_path.exists():
        print(f"Error: {md_path} not found")
        return

    parse_markdown_to_docx(md_path, docx_path)
    print(f"\nGenerated: {docx_path}")
    print(f"File size: {docx_path.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    main()
