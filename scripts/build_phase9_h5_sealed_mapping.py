#!/usr/bin/env python3
"""Build sealed answer-slot to retrieval mapping after review completion."""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OLD_PLAN = ROOT / "runs/phase8_r4_improvements/generation_r4_freeze/request_plan.jsonl"
NEW_PLAN = ROOT / "runs/phase9_h5_followup/generation_freeze/request_plan.jsonl"
OLD_INDEX = ROOT / "runs/phase8_r4_improvements/generation_r4_v3_full/combined_100_record_index.jsonl"
METRICS = ROOT / "runs/phase8_r4_improvements/evaluation_r4/per_query.json"


def rows(path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output-dir", type=Path, default=ROOT / "runs/phase9_h5_followup/analysis_freeze")
    args = parser.parse_args()
    if args.output_dir.exists() and any(args.output_dir.iterdir()):
        raise RuntimeError("analysis freeze exists; refuse overwrite")
    plans = {row["blinded_request_id"]: row for path in (OLD_PLAN, NEW_PLAN) for row in rows(path)}
    sources = {row["blinded_request_id"]: row["source"] for row in rows(OLD_INDEX)}
    metric_doc = json.loads(METRICS.read_text(encoding="utf-8"))
    metrics = {(system, row["query_id"]): row["metrics"] for system, panel in metric_doc.items() for row in panel}
    wb = load_workbook(args.workbook, read_only=True, data_only=True)
    ws = wb["Reviewer A"]
    header = [cell.value for cell in ws[1]]
    workbook_rows = [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)]
    output = []
    for row in workbook_rows:
        plan = plans.get(row["blinded_request_id"])
        if not plan:
            raise RuntimeError(f"unmapped blinded request: {row['blinded_request_id']}")
        source = sources.get(row["blinded_request_id"], "phase9_v3_new")
        prompt_version = "v3_semantic_contract" if source in {"v3_full", "v3_trace", "phase9_v3_new"} else "legacy_v1_v2_contract"
        output.append({
            "answer_slot_id": row["answer_slot_id"], "blinded_request_id": row["blinded_request_id"],
            "query_id": plan["query_id"], "system_id": plan["system_id"], "category": plan["category"],
            "mrr@10": metrics[(plan["system_id"], plan["query_id"])]["mrr@10"],
            "ndcg@10": metrics[(plan["system_id"], plan["query_id"])]["ndcg@10"],
            "recall@10": metrics[(plan["system_id"], plan["query_id"])]["recall@10"],
            "prompt_version": prompt_version, "generation_source": source,
        })
    if len(output) != 150 or len({row["answer_slot_id"] for row in output}) != 150 or len({row["query_id"] for row in output}) != 30:
        raise RuntimeError("sealed mapping coverage invalid")
    args.output_dir.mkdir(parents=True)
    mapping = args.output_dir / "sealed_mapping_150.jsonl"
    mapping.write_text("".join(json.dumps(row, sort_keys=True) + "\n" for row in output), encoding="utf-8")
    manifest = {
        "status": "frozen", "answer_n": 150, "question_n": 30,
        "workbook_sha256": sha(args.workbook), "mapping_sha256": sha(mapping),
        "inputs": {str(path.relative_to(ROOT)): sha(path) for path in (OLD_PLAN, NEW_PLAN, OLD_INDEX, METRICS)},
        "prompt_version_counts": {name: sum(row["prompt_version"] == name for row in output) for name in {row["prompt_version"] for row in output}},
    }
    (args.output_dir / "manifest.json").write_text(json.dumps(manifest, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(manifest, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
