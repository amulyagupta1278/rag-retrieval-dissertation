#!/usr/bin/env python3
"""Validate completed Phase 9 H5 two-reviewer workbook before analysis."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook


REVIEW_SHEETS = ("Reviewer A", "Reviewer B")
EXPECTED_ROWS = 150


def as_bool(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str) and value.strip().lower() in {"true", "false"}:
        return value.strip().lower() == "true"
    raise ValueError(f"invalid Boolean value: {value!r}")


def records(ws):
    header = [cell.value for cell in ws[1]]
    return [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def validate(path: Path, require_complete: bool = True) -> dict:
    wb = load_workbook(path, data_only=False, read_only=True)
    missing = [name for name in (*REVIEW_SHEETS, "Adjudication") if name not in wb.sheetnames]
    if missing:
        raise ValueError(f"missing sheets: {missing}")
    panels = {name: records(wb[name]) for name in REVIEW_SHEETS}
    if any(len(rows) != EXPECTED_ROWS for rows in panels.values()):
        raise ValueError("reviewer panels must contain exactly 150 rows")
    ids = [[row["answer_slot_id"] for row in panels[name]] for name in REVIEW_SHEETS]
    if len(set(ids[0])) != EXPECTED_ROWS or ids[0] != ids[1]:
        raise ValueError("answer IDs are incomplete, duplicated, or misaligned")
    protected = ("answer_slot_id", "protected_status", "blinded_request_id", "question", "reference_answer", "evidence_E01", "evidence_E02", "evidence_E03", "generated_answer", "model_abstained", "model_abstention_reason", "cited_evidence_ids")
    for i in range(EXPECTED_ROWS):
        if any(panels["Reviewer A"][i][key] != panels["Reviewer B"][i][key] for key in protected):
            raise ValueError(f"protected-field mismatch at row {i + 2}")
    pending = sum(row["protected_status"] != "READY_FROZEN" for row in panels["Reviewer A"])
    if require_complete and pending:
        raise ValueError(f"{pending} rows are not READY_FROZEN")
    for sheet_name, rows in panels.items():
        for i, row in enumerate(rows, 2):
            if not require_complete and row["protected_status"] != "READY_FROZEN":
                continue
            status = row["review_status"]
            if status not in {"COMPLETE", "NO_VERIFIABLE_CLAIMS"}:
                raise ValueError(f"{sheet_name}!U{i} incomplete")
            if status == "COMPLETE":
                counts = [row[key] for key in ("total_verifiable_claims", "fully_supported_claims", "partially_supported_claims", "unsupported_claims", "contradicted_claims")]
                if any(type(value) not in (int, float) or value < 0 or int(value) != value for value in counts):
                    raise ValueError(f"{sheet_name} row {i} has invalid claim counts")
                if counts[0] == 0 and not as_bool(row["model_abstained"]):
                    raise ValueError(f"{sheet_name} row {i} has zero claims without abstention")
                if sum(counts[1:]) != counts[0]:
                    raise ValueError(f"{sheet_name} row {i} claim counts do not reconcile")
    adjudicated = records(wb["Adjudication"])
    if len(adjudicated) != EXPECTED_ROWS:
        raise ValueError("adjudication panel must contain exactly 150 rows")
    if require_complete:
        for i, row in enumerate(adjudicated, 2):
            if row["adjudication_status"] not in {"COMPLETE", "NO_VERIFIABLE_CLAIMS", "NOT_REQUIRED"}:
                raise ValueError(f"Adjudication!O{i} incomplete")
            if row["adjudication_status"] in {"COMPLETE", "NOT_REQUIRED"}:
                values = [row[key] for key in ("final_total_claims", "final_fully_supported", "final_partially_supported", "final_unsupported", "final_contradicted")]
                if any(type(value) not in (int, float) or value < 0 or int(value) != value for value in values):
                    raise ValueError(f"Adjudication row {i} has invalid final counts")
                if sum(values[1:]) != values[0]:
                    raise ValueError(f"Adjudication row {i} counts do not reconcile")
    return {"answer_n": EXPECTED_ROWS, "pending_protected_n": pending, "status": "valid_complete" if not pending and require_complete else "valid_draft"}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--allow-draft", action="store_true")
    args = parser.parse_args()
    print(json.dumps(validate(args.workbook, require_complete=not args.allow_draft), indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
