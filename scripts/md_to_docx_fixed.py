#!/usr/bin/env python3
"""
Fixed Markdown to DOCX Converter with Proper Styling
=====================================================
Correctly applies heading styles, inline formatting, and professional layout.
"""

import re
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_cell_background(cell, fill):
    """Set cell background color."""
    shading_elm = OxmlElement("w:shd")
    shading_elm.set(qn("w:fill"), fill)
    cell._element.get_or_add_tcPr().append(shading_elm)


def add_rich_paragraph(doc, text: str, style: str = "Normal"):
    """Add paragraph with proper inline formatting support."""
    p = doc.add_paragraph(style=style)

    # Pattern: **bold**, *italic*, `code`, [link](url)
    pattern = r"\*\*(.+?)\*\*|\*(.+?)\*|`(.+?)`|\[(.+?)\]\((.+?)\)|__(.+?)__"

    last_end = 0
    for match in re.finditer(pattern, text):
        # Add plain text before match
        if match.start() > last_end:
            p.add_run(text[last_end : match.start()])

        # Handle match
        if match.group(1):  # **bold**
            run = p.add_run(match.group(1))
            run.bold = True
        elif match.group(2):  # *italic*
            run = p.add_run(match.group(2))
            run.italic = True
        elif match.group(3):  # `code`
            run = p.add_run(match.group(3))
            run.font.name = "Courier New"
            run.font.size = Pt(10)
        elif match.group(4) and match.group(5):  # [link](url)
            run = p.add_run(match.group(4))
            run.underline = True
            run.font.color.rgb = RGBColor(0, 0, 255)
        elif match.group(6):  # __underline__
            run = p.add_run(match.group(6))
            run.underline = True

        last_end = match.end()

    # Add remaining text
    if last_end < len(text):
        p.add_run(text[last_end:])

    return p


def parse_table(table_lines: list) -> tuple:
    """Parse markdown table and return (headers, rows)."""
    if not table_lines:
        return [], []

    headers = [cell.strip() for cell in table_lines[0].split("|")[1:-1]]
    rows = []

    for table_line in table_lines[1:]:
        if "-" in table_line and "|" in table_line:
            continue  # Skip separator line
        cells = [cell.strip() for cell in table_line.split("|")[1:-1]]
        if cells and len(cells) == len(headers):
            rows.append(cells)

    return headers, rows


def convert_markdown_to_docx(md_path: Path, docx_path: Path) -> None:
    """Convert Markdown to DOCX with proper formatting."""
    doc = Document()

    # Configure default styles
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    with md_path.open("r", encoding="utf-8") as fh:
        lines = fh.readlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Empty line
        if not line.strip():
            i += 1
            continue

        # Heading 1: # text
        if line.startswith("# ") and not line.startswith("## "):
            text = line[2:].strip()
            p = doc.add_heading(text, level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1

        # Heading 2: ## text
        elif line.startswith("## ") and not line.startswith("### "):
            text = line[3:].strip()
            p = doc.add_heading(text, level=2)
            i += 1

        # Heading 3: ### text
        elif line.startswith("### ") and not line.startswith("#### "):
            text = line[4:].strip()
            p = doc.add_heading(text, level=3)
            i += 1

        # Heading 4: #### text
        elif line.startswith("#### "):
            text = line[5:].strip()
            p = doc.add_heading(text, level=4)
            i += 1

        # Horizontal rule: ---
        elif line.strip() == "---":
            doc.add_paragraph()
            i += 1

        # Table: | col | col |
        elif "|" in line:
            table_lines = [line]
            i += 1

            # Collect table rows
            while i < len(lines) and "|" in lines[i]:
                table_lines.append(lines[i].rstrip())
                i += 1

            headers, rows = parse_table(table_lines)

            if headers and rows:
                table = doc.add_table(rows=1 + len(rows), cols=len(headers))
                table.style = "Light Grid Accent 1"

                # Header row
                hdr_cells = table.rows[0].cells
                for idx, header in enumerate(headers):
                    hdr_cells[idx].text = header
                    # Style header
                    for paragraph in hdr_cells[idx].paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                    set_cell_background(hdr_cells[idx], "D3D3D3")

                # Data rows
                for row_idx, row_data in enumerate(rows, 1):
                    cells = table.rows[row_idx].cells
                    for col_idx, cell_text in enumerate(row_data):
                        cells[col_idx].text = cell_text

        # Bullet list: - text
        elif line.strip().startswith("- "):
            text = line.strip()[2:].strip()
            add_rich_paragraph(doc, text, style="List Bullet")
            i += 1

        # Numbered list: 1. text
        elif re.match(r"^\s*\d+\.\s+", line):
            match = re.match(r"^\s*\d+\.\s+(.*)", line)
            if match:
                text = match.group(1).strip()
                add_rich_paragraph(doc, text, style="List Number")
            i += 1

        # Code block: ```
        elif line.strip().startswith("```"):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip())
                i += 1
            if i < len(lines):
                i += 1  # Skip closing ```

            code_text = "\n".join(code_lines).strip()
            if code_text:
                p = doc.add_paragraph(code_text, style="No Spacing")
                for run in p.runs:
                    run.font.name = "Courier New"
                    run.font.size = Pt(9)

        # Regular paragraph
        else:
            text = line.strip()
            if text:
                add_rich_paragraph(doc, text)
            i += 1

    # Set margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    doc.save(docx_path)


def main():
    base_path = Path(__file__).parent.parent
    md_path = base_path / "FINAL_COMPILED_MIDSEM_REPORT.md"
    docx_path = base_path / "FINAL_COMPILED_MIDSEM_REPORT.docx"

    if not md_path.exists():
        print(f"✗ Error: {md_path} not found")
        return

    convert_markdown_to_docx(md_path, docx_path)

    print(f"✓ Generated: {docx_path}")
    print(f"✓ File size: {docx_path.stat().st_size / 1024:.1f} KB")
    print("\n✓ Formatting applied:")
    print("  • Heading levels 1–4 (proper styles)")
    print("  • Inline: **bold**, *italic*, `code`")
    print("  • Markdown tables with header shading")
    print("  • Bullet and numbered lists")
    print("  • Professional margins (1 inch)")


if __name__ == "__main__":
    main()
