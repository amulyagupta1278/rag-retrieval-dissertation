#!/usr/bin/env python3
"""Validate completed H5c workbook, unseal mapping, run preregistered analysis."""
from __future__ import annotations
import argparse, json
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from scipy.stats import binomtest
SEED=51025; BOOT=10000; N=60
def rows(ws):
 h=[c.value for c in ws[1]]; return [dict(zip(h,r)) for r in ws.iter_rows(min_row=2,values_only=True)]
def as_bool(v):
 if isinstance(v,bool): return v
 if isinstance(v,(int,float)) and v in (0,1): return bool(v)
 if isinstance(v,str) and v.strip().lower() in {'true','false'}: return v.strip().lower()=='true'
 raise ValueError(f'invalid boolean {v!r}')
def validate(path):
 wb=load_workbook(path,read_only=True,data_only=True)
 for name in ('Reviewer A','Reviewer B','Adjudication'):
  if name not in wb.sheetnames: raise ValueError(f'missing sheet {name}')
 panels={x:rows(wb[x]) for x in ('Reviewer A','Reviewer B')}; adj=rows(wb['Adjudication'])
 if any(len(x)!=N for x in [*panels.values(),adj]): raise ValueError('expected 60 rows per scoring sheet')
 ids=[x['answer_slot_id'] for x in panels['Reviewer A']]
 if len(set(ids))!=N or ids!=[x['answer_slot_id'] for x in panels['Reviewer B']]: raise ValueError('slot coverage differs')
 protected=('answer_slot_id','protected_status','blinded_request_id','question','evidence_E01','evidence_E02','evidence_E03','generated_answer','model_abstained','model_abstention_reason','cited_evidence_ids')
 for i in range(N):
  if panels['Reviewer A'][i]['protected_status']!='READY_FROZEN': raise ValueError(f'row {i+2} not READY_FROZEN')
  if any(panels['Reviewer A'][i][k]!=panels['Reviewer B'][i][k] for k in protected): raise ValueError(f'protected mismatch row {i+2}')
  for name in panels:
   r=panels[name][i]; status=r['review_status']
   if status not in {'COMPLETE','NO_VERIFIABLE_CLAIMS'}: raise ValueError(f'{name} row {i+2} incomplete')
   vals=[r[k] or 0 for k in ('total_verifiable_claims','fully_supported_claims','partially_supported_claims','unsupported_claims','contradicted_claims')]
   if any(type(v) not in (int,float) or v<0 or int(v)!=v for v in vals) or sum(vals[1:])!=vals[0]: raise ValueError(f'{name} row {i+2} counts invalid')
   if vals[0]==0 and not as_bool(r['model_abstained']): raise ValueError(f'{name} row {i+2} zero claims without abstention')
  ar=adj[i]
  if ar['adjudication_status'] not in {'COMPLETE','NO_VERIFIABLE_CLAIMS','NOT_REQUIRED'}: raise ValueError(f'adjudication row {i+2} incomplete')
  vals=[ar[k] or 0 for k in ('final_total_claims','final_fully_supported','final_partially_supported','final_unsupported','final_contradicted')]
  if any(type(v) not in (int,float) or v<0 or int(v)!=v for v in vals) or sum(vals[1:])!=vals[0]: raise ValueError(f'adjudication row {i+2} counts invalid')
 return panels['Reviewer A'],adj
def main():
 ap=argparse.ArgumentParser(); ap.add_argument('workbook',type=Path); ap.add_argument('mapping',type=Path); ap.add_argument('--output',type=Path,required=True); args=ap.parse_args()
 protected,adj=validate(args.workbook); mapping=[json.loads(x) for x in args.mapping.read_text().splitlines() if x]
 if len(mapping)!=N or len({x['blinded_request_id'] for x in mapping})!=N: raise ValueError('mapping coverage invalid')
 byblind={x['blinded_request_id']:(p,a) for x,p,a in zip(protected,protected,adj)}; obs=[]
 for m in mapping:
  p,a=byblind[m['blinded_request_id']]; total=a['final_total_claims'] or 0; faith=((a['final_fully_supported'] or 0)+.5*(a['final_partially_supported'] or 0))/total if total else None; answered=not as_bool(p['model_abstained']); success=bool(answered and faith is not None and faith>=.90); obs.append({**m,'answered':answered,'faithfulness':faith,'grounded_success':success})
 paired={q:{x['condition']:x for x in obs if x['query_id']==q} for q in sorted({x['query_id'] for x in obs})}; comp=[int(v['complete_top3']['grounded_success']) for v in paired.values()]; irr=[int(v['irrelevant']['grounded_success']) for v in paired.values()]; b=sum(x==1 and y==0 for x,y in zip(comp,irr)); c=sum(x==0 and y==1 for x,y in zip(comp,irr)); discord=b+c; p=1.0 if not discord else float(binomtest(min(b,c),discord,.5,alternative='two-sided').pvalue); diff=float(np.mean(comp)-np.mean(irr)); rng=np.random.default_rng(SEED); draws=[]
 for _ in range(BOOT):
  ix=rng.integers(0,len(comp),len(comp)); draws.append(float(np.mean([comp[i]-irr[i] for i in ix])))
 ci=[float(np.percentile(draws,2.5)),float(np.percentile(draws,97.5))]; supported=bool(np.mean(comp)>np.mean(irr) and diff>=.30 and p<.05 and ci[0]>0)
 cond={k:{'n':sum(x['condition']==k for x in obs),'answer_rate':float(np.mean([x['answered'] for x in obs if x['condition']==k])),'grounded_success_rate':float(np.mean([x['grounded_success'] for x in obs if x['condition']==k]))} for k in sorted({x['condition'] for x in obs})}
 result={'hypothesis':'H5c','decision':'supported' if supported else 'not_supported','confirmatory':True,'question_n':len(paired),'response_n':len(obs),'primary_contrast':'complete_top3 versus irrelevant','complete_success_rate':float(np.mean(comp)),'irrelevant_success_rate':float(np.mean(irr)),'paired_difference':diff,'bootstrap_ci95':ci,'mcnemar_b_complete_only':b,'mcnemar_c_irrelevant_only':c,'mcnemar_exact_two_sided_p':p,'bootstrap_samples':BOOT,'seed':SEED,'condition_summary':cond,'original_H5_status_unchanged':'not_estimable'}
 args.output.parent.mkdir(parents=True,exist_ok=True); args.output.write_text(json.dumps(result,indent=2,sort_keys=True)+'\n'); print(json.dumps(result,indent=2,sort_keys=True))
if __name__=='__main__': main()
