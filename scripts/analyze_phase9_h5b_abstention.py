#!/usr/bin/env python3
"""Exploratory H5b: retrieval quality versus answer coverage/abstention."""

from __future__ import annotations

import argparse
import json
import math
from collections import defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.stats import spearmanr

SEED = 42
SAMPLES = 10_000
METRICS = ("mrr@10", "ndcg@10", "recall@10")


def as_bool(value):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    if isinstance(value, str) and value.strip().lower() in {"true", "false"}:
        return value.strip().lower() == "true"
    raise ValueError(f"invalid Boolean value: {value!r}")


def jsonl(path):
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line]


def workbook_rows(path, sheet):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet]
    header = [cell.value for cell in ws[1]]
    return [dict(zip(header, row)) for row in ws.iter_rows(min_row=2, values_only=True)]


def stat(rows, metric):
    x = np.asarray([row[metric] for row in rows], dtype=float)
    answered = np.asarray([row["answered"] for row in rows], dtype=float)
    rho = spearmanr(x, answered).statistic
    answered_mean = float(x[answered == 1].mean())
    abstained_mean = float(x[answered == 0].mean())
    q25, q75 = np.quantile(x, [0.25, 0.75])
    low = answered[x <= q25]
    high = answered[x >= q75]
    return {
        "spearman_rho_answered": None if math.isnan(rho) else float(rho),
        "mean_metric_answered": answered_mean,
        "mean_metric_abstained": abstained_mean,
        "mean_difference_answered_minus_abstained": answered_mean - abstained_mean,
        "low_quartile_answer_rate": float(low.mean()),
        "high_quartile_answer_rate": float(high.mean()),
        "high_minus_low_answer_rate": float(high.mean() - low.mean()),
    }


def bootstrap(rows, metric):
    grouped = defaultdict(list)
    for row in rows:
        grouped[row["query_id"]].append(row)
    qids = sorted(grouped)
    rng = np.random.default_rng(SEED)
    draws = defaultdict(list)
    for _ in range(SAMPLES):
        sample = [row for qid in rng.choice(qids, len(qids), replace=True) for row in grouped[qid]]
        result = stat(sample, metric)
        for key in ("spearman_rho_answered", "mean_difference_answered_minus_abstained", "high_minus_low_answer_rate"):
            value = result[key]
            if value is not None and not math.isnan(value):
                draws[key].append(value)
    return {key + "_ci95": [float(np.percentile(values, 2.5)), float(np.percentile(values, 97.5))] for key, values in draws.items()} | {key + "_bootstrap_valid_n": len(values) for key, values in draws.items()}


def rates(rows, key):
    grouped = defaultdict(list)
    for row in rows:
        grouped[row[key]].append(row["answered"])
    return {name: {"answer_n": int(sum(values)), "n": len(values), "answer_rate": sum(values) / len(values), "abstention_rate": 1 - sum(values) / len(values)} for name, values in sorted(grouped.items())}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("mapping", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    a = workbook_rows(args.workbook, "Reviewer A")
    b = workbook_rows(args.workbook, "Reviewer B")
    if len(a) != 150 or len(b) != 150:
        raise ValueError("expected two 150-row panels")
    if any(as_bool(a[i]["model_abstained"]) != as_bool(b[i]["model_abstained"]) for i in range(150)):
        raise ValueError("reviewer protected abstention fields differ")
    mapping = {row["answer_slot_id"]: row for row in jsonl(args.mapping)}
    panel = []
    for row in a:
        item = mapping[row["answer_slot_id"]]
        panel.append({**item, "answered": 0 if as_bool(row["model_abstained"]) else 1})
    analyses = {metric: stat(panel, metric) | bootstrap(panel, metric) for metric in METRICS}
    result = {
        "status": "complete_exploratory_post_hoc",
        "hypothesis": "H5b: retrieval quality affects answer coverage/abstention more strongly than faithfulness",
        "confirmatory_use": "prohibited; H5b was formulated after observing H5 ceiling effect",
        "answer_n": int(sum(row["answered"] for row in panel)),
        "abstention_n": int(sum(not row["answered"] for row in panel)),
        "record_n": len(panel),
        "question_n": len({row["query_id"] for row in panel}),
        "bootstrap": {"samples": SAMPLES, "seed": SEED, "unit": "whole question preserving five-system panel"},
        "metric_associations": analyses,
        "by_system": rates(panel, "system_id"),
        "by_category": rates(panel, "category"),
        "by_prompt_version": rates(panel, "prompt_version"),
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
